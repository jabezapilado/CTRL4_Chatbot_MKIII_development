"""
Retrieval-Augmented Generation (RAG) Service

Builds and manages the Guidance Office
knowledge retrieval system.

Responsibilities

- Document Processing
- Text Chunking
- Embedding Generation
- FAISS Index Management
- Semantic Retrieval

CTRL4 Chatbot MK III

Authors:
- Apilado, Jabez Timothy E.
- Quilantang, Grant Mihkael D.
- Lanix, Iligan
- Wylengco, Teyshaun Zell
"""

from __future__ import annotations

import json
import hashlib
import logging
import re
import shutil

from dataclasses import dataclass
from pathlib import Path
from typing import Any
from typing import Final

SUPPORTED_DOCUMENT_TYPES: Final[frozenset[str]] = frozenset({
    ".pdf",
    ".docx",
    ".json",
    ".txt",
    ".md",
})

from ..config import Config


logger = logging.getLogger(__name__)

INDEX_MANIFEST_VERSION: Final = 1
INDEX_MANIFEST_FILENAME: Final = "manifest.json"
LEGACY_INDEX_ARCHIVE_NAME: Final = "legacy_pre_manifest"

@dataclass
class RetrievedDocument:
    text: str
    source: str
    score: float


@dataclass(frozen=True)
class RAGIndexBuildReport:
    """Safe build metadata; it never retains knowledge-record contents."""

    discovered_sources: tuple[str, ...]
    indexed_sources: tuple[str, ...]
    skipped_sources: tuple[dict[str, str], ...]
    chunk_count: int
    source_fingerprint: str

def normalize_text(text: str) -> str:
    lowered = str(text).strip().lower()
    lowered = re.sub(r"\s+", " ", lowered)
    return lowered

def split_text(text: str, chunk_size: int, overlap: int) -> list[str]:
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    if not text:
        return []
    if len(text) <= chunk_size:
        return [text]

    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + chunk_size, len(text))
        chunk = text[start:end]
        if end < len(text):
            split_at = max(chunk.rfind("\n\n"), chunk.rfind(". "), chunk.rfind("? "), chunk.rfind("! "))
            if split_at > chunk_size // 3:
                end = start + split_at + 1
                chunk = text[start:end]
        chunks.append(chunk.strip())
        if end >= len(text):
            break
        start = max(0, end - overlap)

    return [c for c in chunks if c]

def _extract_pdf_text(file_path: Path) -> str:
    try:
        from pypdf import PdfReader  # type: ignore
    except Exception:
        return ""

    pages: list[str] = []
    reader = PdfReader(str(file_path))
    for page in reader.pages:
        pages.append(page.extract_text() or "")
    return "\n".join(pages)


def _extract_docx_text(file_path: Path) -> str:
    try:
        from docx import Document  # type: ignore
    except Exception:
        return ""

    document = Document(str(file_path))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _extract_json_text(file_path: Path) -> str:
    try:
        content = json.loads(file_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError):
        return ""

    return json.dumps(content, ensure_ascii=False, indent=2)


def read_document_text(file_path: Path) -> str:
    suffix = file_path.suffix.lower()
    if suffix == ".pdf":
        return _extract_pdf_text(file_path)
    if suffix == ".docx":
        return _extract_docx_text(file_path)
    if suffix == ".json":
        return _extract_json_text(file_path)
    if suffix in {".txt", ".md"}:
        return file_path.read_text(encoding="utf-8", errors="ignore")
    return ""


def _sha256_file(file_path: Path) -> str:
    digest = hashlib.sha256()
    with file_path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()

class RAGService:
    
    def __init__(self):

        self.config = Config()

        self.docs_dir = Path(self.config.RAG_DOCS_DIR)

        self.index_dir = Path(self.config.RAG_INDEX_DIR)

        self.index_file = self.index_dir / "knowledge.faiss"

        self.metadata_file = self.index_dir / "metadata.json"

        self.manifest_file = self.index_dir / INDEX_MANIFEST_FILENAME

        self.embedding_model_name = self.config.RAG_EMBEDDING_MODEL

        self.embedder: Any | None = None
        self.faiss: Any | None = None
        self.index: Any | None = None
        self.metadata: list[dict[str, Any]] = []

        self.last_build_report: RAGIndexBuildReport | None = None

        self.initialization_error: str | None = None

        self._initialize_runtime()
    
    def _build_embedding_input(self, text: str, is_query: bool) -> str:
        if "e5" in self.embedding_model_name.lower():
            return f"query: {text}" if is_query else f"passage: {text}"
        return text

    def _embed(
        self,
        texts: list[str],
        is_query: bool,
    ) -> Any:

        if self.embedder is None:
            raise RuntimeError(
                "Embedding model is not available."
            )

        prepared = [
            self._build_embedding_input(text, is_query=is_query)
            for text in texts
        ]

        return self.embedder.encode(
            prepared,
            normalize_embeddings=True,
            convert_to_numpy=True,
        )

    def _initialize_runtime(self) -> None:
        try:
            import faiss  # type: ignore
            from sentence_transformers import SentenceTransformer  # type: ignore
        except Exception:
            self.initialization_error = (
                "RAG dependencies are missing. Install sentence-transformers and faiss-cpu."
            )
            return

        self.faiss = faiss
        self.embedder = SentenceTransformer(self.embedding_model_name)

        if self.index_file.exists() and self.metadata_file.exists():
            is_current, reason = self._index_is_current()
            if is_current:
                self._load_index()
                return

            logger.warning("RAG index is stale (reason=%s).", reason)
            if self.config.RAG_AUTO_BUILD_ON_START:
                self.build_index()
                return

            self.initialization_error = (
                "RAG index is stale. Rebuild it from the configured knowledge "
                "base before starting with CHATBOT_RAG_AUTO_BUILD_ON_START=false."
            )
            return

        if self.config.RAG_AUTO_BUILD_ON_START:
            self.build_index()

    def _load_index(self) -> None:
        if not self.faiss:
            return
        self.index = self.faiss.read_index(str(self.index_file))
        self.metadata: list[dict[str, Any]] = json.loads(
            self.metadata_file.read_text(encoding="utf-8")
        )

    def _discover_source_paths(self) -> list[Path]:
        if not self.docs_dir.is_dir():
            return []

        return sorted(
            path
            for path in self.docs_dir.rglob("*")
            if path.is_file()
            and not path.name.startswith("._")
            and path.suffix.lower() in SUPPORTED_DOCUMENT_TYPES
        )

    def _source_manifest(self, source_paths: list[Path]) -> dict[str, Any]:
        sources = [
            {
                "path": str(path.relative_to(self.docs_dir)),
                "sha256": _sha256_file(path),
            }
            for path in source_paths
        ]
        source_definition = {
            "version": INDEX_MANIFEST_VERSION,
            "embedding_model": self.embedding_model_name,
            "chunk_size": self.config.RAG_CHUNK_SIZE,
            "chunk_overlap": self.config.RAG_CHUNK_OVERLAP,
            "sources": sources,
        }
        serialized = json.dumps(
            source_definition,
            ensure_ascii=False,
            sort_keys=True,
            separators=(",", ":"),
        )
        return {
            **source_definition,
            "source_fingerprint": hashlib.sha256(
                serialized.encode("utf-8")
            ).hexdigest(),
        }

    def _index_is_current(self) -> tuple[bool, str]:
        if not self.manifest_file.is_file():
            return False, "manifest_missing"

        try:
            manifest = json.loads(self.manifest_file.read_text(encoding="utf-8"))
            source_manifest = self._source_manifest(self._discover_source_paths())
            metadata = json.loads(self.metadata_file.read_text(encoding="utf-8"))
        except (OSError, ValueError, TypeError):
            return False, "manifest_or_metadata_invalid"

        if manifest.get("source_fingerprint") != source_manifest["source_fingerprint"]:
            return False, "source_fingerprint_mismatch"

        indexed_sources = manifest.get("indexed_sources")
        if not isinstance(indexed_sources, list):
            return False, "indexed_sources_invalid"

        metadata_sources = {
            str(item.get("source", ""))
            for item in metadata
            if isinstance(item, dict)
        }
        if metadata_sources != set(indexed_sources):
            return False, "metadata_source_mismatch"

        if manifest.get("chunk_count") != len(metadata):
            return False, "chunk_count_mismatch"

        return True, "current"

    def _backup_legacy_index(self) -> None:
        """Preserve a pre-manifest index once without making it runtime input."""
        if self.manifest_file.exists() or not self.index_file.exists():
            return

        archive_dir = self.index_dir / "archive" / LEGACY_INDEX_ARCHIVE_NAME
        if archive_dir.exists():
            return

        archive_dir.mkdir(parents=True, exist_ok=False)
        shutil.copy2(self.index_file, archive_dir / self.index_file.name)
        if self.metadata_file.exists():
            shutil.copy2(self.metadata_file, archive_dir / self.metadata_file.name)

    @staticmethod
    def _write_json(path: Path, value: object) -> None:
        temporary_path = path.with_suffix(path.suffix + ".tmp")
        temporary_path.write_text(
            json.dumps(value, ensure_ascii=False, indent=2),
            encoding="utf-8",
        )
        temporary_path.replace(path)

    def build_index(self) -> int:
        if not self.embedder or not self.faiss:
            return 0

        self.index_dir.mkdir(parents=True, exist_ok=True)
        self.docs_dir.mkdir(parents=True, exist_ok=True)

        self._backup_legacy_index()

        doc_paths = self._discover_source_paths()
        source_manifest = self._source_manifest(doc_paths)

        chunks: list[dict[str, Any]] = []
        indexed_sources: list[str] = []
        skipped_sources: list[dict[str, str]] = []
        for path in doc_paths:
            relative_path = str(path.relative_to(self.docs_dir))
            try:
                raw = read_document_text(path)
            except (OSError, ValueError):
                skipped_sources.append(
                    {"source": relative_path, "reason": "unreadable"}
                )
                continue

            if not normalize_text(raw):
                skipped_sources.append(
                    {"source": relative_path, "reason": "empty_or_invalid"}
                )
                continue

            source_chunk_count = 0
            for idx, chunk in enumerate(
                split_text(
                    raw,
                    chunk_size=self.config.RAG_CHUNK_SIZE,
                    overlap=self.config.RAG_CHUNK_OVERLAP,
                )
            ):
                chunks.append(
                    {
                        "text": chunk,
                        "source": relative_path,
                        "chunk_id": idx,
                    }
                )
                source_chunk_count += 1

            if source_chunk_count:
                indexed_sources.append(relative_path)
            else:
                skipped_sources.append(
                    {"source": relative_path, "reason": "no_indexable_chunks"}
                )

        if not chunks:
            self.index = None
            self.metadata = []
            self.last_build_report = RAGIndexBuildReport(
                discovered_sources=tuple(
                    item["path"] for item in source_manifest["sources"]
                ),
                indexed_sources=(),
                skipped_sources=tuple(skipped_sources),
                chunk_count=0,
                source_fingerprint=source_manifest["source_fingerprint"],
            )
            self.initialization_error = "RAG index build found no indexable knowledge records."
            return 0

        vectors = self._embed([item["text"] for item in chunks], is_query=False)
        dim = int(vectors.shape[1])
        index = self.faiss.IndexFlatIP(dim)
        index.add(vectors)

        temporary_index_file = self.index_file.with_suffix(".faiss.tmp")
        self.faiss.write_index(index, str(temporary_index_file))
        temporary_index_file.replace(self.index_file)

        index_manifest = {
            **source_manifest,
            "indexed_sources": indexed_sources,
            "skipped_sources": skipped_sources,
            "chunk_count": len(chunks),
        }
        self._write_json(self.metadata_file, chunks)
        self._write_json(self.manifest_file, index_manifest)

        self.index = index
        self.metadata = chunks
        self.initialization_error = None
        self.last_build_report = RAGIndexBuildReport(
            discovered_sources=tuple(item["path"] for item in source_manifest["sources"]),
            indexed_sources=tuple(indexed_sources),
            skipped_sources=tuple(skipped_sources),
            chunk_count=len(chunks),
            source_fingerprint=source_manifest["source_fingerprint"],
        )
        logger.info(
            "RAG index built (sources=%s indexed=%s skipped=%s chunks=%s fingerprint=%s).",
            len(doc_paths),
            len(indexed_sources),
            len(skipped_sources),
            len(chunks),
            source_manifest["source_fingerprint"][:12],
        )
        return len(chunks)

    def retrieve(self, query: str) -> list[RetrievedDocument]:

        if not self.embedder or not self.faiss or self.index is None or not self.metadata:
            return []

        query_vector = self._embed([query], is_query=True)

        k = max(1, self.config.RAG_TOP_K)

        scores, indices = self.index.search(query_vector, k)

        results: list[RetrievedDocument] = []
        seen_chunks: set[str] = set()

        for score, idx in zip(scores[0], indices[0]):

            if idx < 0:
                continue

            if float(score) < self.config.RAG_MIN_SCORE:
                continue

            meta = self.metadata[int(idx)]

            text = str(meta.get("text", ""))

            if text in seen_chunks:
                continue

            seen_chunks.add(text)

            results.append(
                RetrievedDocument(
                    text=text,
                    source=str(meta.get("source", "unknown")),
                    score=float(score),
                )
            )

        return results
    
    def reload_index(self) -> int:

        if not self.index_file.exists():
            return 0

        self._load_index()

        return len(self.metadata)
    
    def stats(self) -> dict[str, object]:

        return {
            "documents": len(self.metadata),
            "embedding_model": self.embedding_model_name,
            "indexed": self.index is not None,
            "documents_directory": str(self.docs_dir),
            "index_directory": str(self.index_dir),
            "source_fingerprint": (
                self.last_build_report.source_fingerprint
                if self.last_build_report
                else None
            ),
        }
        
    @property
    def ready(self) -> bool:
        return (
            self.embedder is not None
            and self.index is not None
            and self.faiss is not None
        )
